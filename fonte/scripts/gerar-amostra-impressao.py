"""EFCdocs v0.6.2 — gera testes/amostras/amostra-impressao.xlsx (dados INVENTADOS) para testar Imprimir e PDF:
datas em coluna estreita, texto longo ao lado de célula cheia, quebra de linha com Alt+Enter e mesclada que atravessa a quebra de página."""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

wb = Workbook()
ws = wb.active
ws.title = 'Entregas'
cab = ['Data', 'Responsável', 'Observação', 'Valor', 'Situação', 'Resumo']
ws.append(cab)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F4C5C'); c.alignment = Alignment(horizontal='center')
nomes = ['Ana Souza', 'Bruno Lima', 'Carla Dias', 'Diego Rocha']
obs = ['Entrega conferida no portão com todos os itens da lista', 'Faltou um item', 'Tudo certo', 'Cliente pediu para remarcar a visita para a semana seguinte']
for i in range(2, 90):
    d = datetime.date(2026, 1, 5) + datetime.timedelta(days=i * 3)
    ws.append([d, nomes[i % 4], obs[i % 4], round(120.5 + i * 7.35, 2), 'OK' if i % 5 else 'Pendente', None])
    ws.cell(i, 1).number_format = 'dd/mm/yyyy'
    ws.cell(i, 4).number_format = '"R$" #,##0.00'
ws.cell(5, 3).value = 'Coletar\namostra\nno ponto 2'
ws.cell(5, 3).alignment = Alignment(wrap_text=False)
ws.row_dimensions[5].height = 45
ws['F2'] = 'Revisão geral do mês'
ws.merge_cells('F2:F70')
ws['F2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws['F2'].fill = PatternFill('solid', fgColor='E8EEF1')
ws.column_dimensions['A'].width = 10.5
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 24
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 11
ws.column_dimensions['F'].width = 16
wb.save('testes/amostras/amostra-impressao.xlsx')
print('ok')
