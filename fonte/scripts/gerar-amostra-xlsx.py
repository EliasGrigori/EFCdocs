"""EFCdocs — gera a planilha de teste do Excel (dados inventados, nada real).
Uso: python3 scripts/gerar-amostra-xlsx.py testes/amostras/amostra-qualidade.xlsx"""
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

saida = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = 'Coletas'
ws.sheet_properties.tabColor = '1F7A46'
cab = ['Data', 'Ponto', 'Cloro (mg/L)', 'pH', 'Situação', 'Mês']
ws.append(cab)
fino = Side(style='thin', color='999999')
for i, t in enumerate(cab, start=1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='0F4C5C')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = Border(top=fino, bottom=fino, left=fino, right=fino)
pontos = ['Entrada', 'Reservatório', 'Rede 1', 'Rede 2', 'Bebedouro']
base = datetime.date(2026, 1, 5)
for k in range(40):
    r = k + 2
    ws.cell(row=r, column=1, value=base + datetime.timedelta(days=7 * (k // 5))).number_format = 'dd/mm/yyyy'
    ws.cell(row=r, column=2, value=pontos[k % 5])
    ws.cell(row=r, column=3, value=round(0.4 + (k * 37 % 17) / 10, 2)).number_format = '0.00'
    ws.cell(row=r, column=4, value=round(6.6 + (k * 13 % 9) / 10, 1))
    ws.cell(row=r, column=5, value=f'=IF(AND(C{r}>=0.2,C{r}<=2),"CONFORME","NÃO CONFORME")')
    ws.cell(row=r, column=6, value=f'=UPPER(TEXT(A{r},"mmmm"))')
# fórmula copiada até bem longe, dando vazio (o EFCdocs corta e avisa)
for r in range(42, 6001):
    ws.cell(row=r, column=5, value=f'=IF(C{r}="","",IF(AND(C{r}>=0.2,C{r}<=2),"CONFORME","NÃO CONFORME"))')
tab = Table(displayName='TAB_COLETAS', ref='A1:F41')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = 'A2'
ws.column_dimensions['B'].width = 18
ws.column_dimensions['E'].width = 16
ws.row_dimensions[1].height = 30
dv = DataValidation(type='list', formula1='"Entrada,Reservatório,Rede 1,Rede 2,Bebedouro"')
ws.add_data_validation(dv); dv.add('B2:B41')
ws.conditional_formatting.add('C2:C41', CellIsRule(operator='greaterThan', formula=['2'], fill=PatternFill('solid', fgColor='F4CCCC')))
ws['B2'].comment = Comment('Ponto de entrada da água', 'Teste')

rs = wb.create_sheet('Resumo')
rs['A1'] = 'Resumo do mês'; rs['A1'].font = Font(bold=True, size=14)
rs.merge_cells('A1:C1')
rs['A1'].alignment = Alignment(horizontal='center')
rs['A3'] = 'Média do cloro'; rs['B3'] = '=AVERAGE(TAB_COLETAS[Cloro (mg/L)])'; rs['B3'].number_format = '0.00'
rs['A4'] = 'Fora do padrão'; rs['B4'] = '=COUNTIF(TAB_COLETAS[Situação],"NÃO CONFORME")'
rs['A5'] = 'Maior pH'; rs['B5'] = '=MAX(Coletas!D2:D41)'
rs['A6'] = 'Cloro da Rede 1'; rs['B6'] = '=SUMIF(Coletas!B2:B41,"Rede 1",Coletas!C2:C41)'
rs['A7'] = 'Procura'; rs['B7'] = '=VLOOKUP("Rede 2",Coletas!B2:C41,2,FALSE)'
rs['A8'] = 'Meta'; rs['B8'] = 0.5
rs['A9'] = 'Acima da meta'; rs['B9'] = '=COUNTIF(Coletas!C2:C41,">"&Meta_Cloro)'
rs['A10'] = 'Moeda'; rs['B10'] = 1234.5; rs['B10'].number_format = '"R$" #,##0.00'
rs['A11'] = 'Porcentagem'; rs['B11'] = 0.257; rs['B11'].number_format = '0.0%'
rs['A12'] = 'Verdadeiro?'; rs['B12'] = True
rs['A13'] = 'Erro'; rs['B13'] = '=1/0'
rs.column_dimensions['A'].width = 22
rs.column_dimensions['D'].hidden = True
rs.row_dimensions[14].hidden = True
rs['A14'] = 'linha oculta'
wb.defined_names['Meta_Cloro'] = DefinedName('Meta_Cloro', attr_text='Resumo!$B$8')
oc = wb.create_sheet('Apoio')
oc.sheet_state = 'hidden'
oc['A1'] = 'aba oculta'
wb.save(saida)
print('ok', saida)
