"""EFCdocs — gera testes/amostras/amostra-formas.xlsx (DADOS INVENTADOS).
Imita um painel feito no Excel com FORMAS: painel escuro de fundo (degradê), título em caixa de texto, grupo de botões de
navegação com link para outras abas (cantos em diagonal, sombra), setas (bloco e linha com ponta), documento de fluxograma,
desenho livre, caixa de texto ligada a uma célula, um grupo com um gráfico dentro e um grupo com uma figura dentro.
Uso: python3 testes/amostras/gerar-amostra-formas.py"""
import io, os, struct, zipfile, zlib
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'amostra-formas.xlsx')
EMU = 9525

wb = Workbook()
ws = wb.active
ws.title = 'Painel'
ws['B40'] = 'Total de não conformidades'
ws['C40'] = '=SUM(Dados!B2:B7)'
dados = wb.create_sheet('Dados')
dados['A1'] = 'Mês'; dados['B1'] = 'Não conformidades'; dados['C1'] = 'Meta'
for i, (m, v) in enumerate([('Jan', 7), ('Fev', 5), ('Mar', 9), ('Abr', 4), ('Mai', 3), ('Jun', 6)], start=2):
    dados.cell(i, 1, m); dados.cell(i, 2, v); dados.cell(i, 3, 5)
acoes = wb.create_sheet('Ações Corretivas')
acoes['A1'] = 'Ação'; acoes['B1'] = 'Responsável'; acoes['C1'] = 'Prazo'
for i, (a, r, p) in enumerate([('Trocar filtro do bebedouro', 'Ana', '10/11'), ('Recalibrar balança', 'Bruno', '15/11'), ('Treinar equipe do packing', 'Carla', '30/11')], start=2):
    acoes.cell(i, 1, a); acoes.cell(i, 2, r); acoes.cell(i, 3, p)
ch = BarChart()
ch.title = 'Não conformidades por mês'
ch.add_data(Reference(dados, min_col=2, min_row=1, max_row=7), titles_from_data=True)
ch.set_categories(Reference(dados, min_col=1, min_row=2, max_row=7))
ws.add_chart(ch, 'E14')
buf = io.BytesIO()
wb.save(buf)


def png(w, h, cor):
    """PNG pequeno de uma cor (sem biblioteca)."""
    linha = b'\x00' + bytes(cor) * w
    bruto = linha * h
    def bloco(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + bloco(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + bloco(b'IDAT', zlib.compress(bruto)) + bloco(b'IEND', b'')


def ancora(c1, r1, c2, r2, dentro, edit=''):
    ea = f' editAs="{edit}"' if edit else ''
    return (f'<xdr:twoCellAnchor{ea}><xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{c2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>{dentro}<xdr:clientData/></xdr:twoCellAnchor>')


def xfrm(x, y, w, h, extra=''):
    return f'<a:xfrm{extra}><a:off x="{x * EMU}" y="{y * EMU}"/><a:ext cx="{w * EMU}" cy="{h * EMU}"/></a:xfrm>'


ESTILO_BOTAO = ('<xdr:style><a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="15000"/></a:schemeClr></a:lnRef><a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef>'
                '<a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef><a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></xdr:style>')


def texto(t, sz=1100, b=False, cor='FFFFFF', algn='ctr', anchor='ctr', extra_bp=''):
    return (f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" rtlCol="0" anchor="{anchor}"{extra_bp}/><a:lstStyle/><a:p><a:pPr algn="{algn}"/>'
            f'<a:r><a:rPr lang="pt-BR" sz="{sz}" b="{1 if b else 0}"><a:solidFill><a:srgbClr val="{cor}"/></a:solidFill></a:rPr><a:t>{t}</a:t></a:r></a:p></xdr:txBody>')


def botao(id_, nome, rid, x, y, w, h, sombra=False):
    ef = ('<a:effectLst><a:outerShdw blurRad="76200" dir="13500000" sy="23000" kx="1200000" algn="br" rotWithShape="0"><a:prstClr val="black"><a:alpha val="20000"/></a:prstClr></a:outerShdw></a:effectLst>' if sombra else '')
    return (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{id_}" name="Botão {nome}"><a:hlinkClick xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"/></xdr:cNvPr><xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="round2DiagRect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="accent6"/></a:solidFill><a:ln w="6350"><a:noFill/></a:ln>{ef}</xdr:spPr>'
            f'{ESTILO_BOTAO}<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="pt-BR" sz="1100" b="1"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill><a:latin typeface="+mn-lt"/></a:rPr><a:t>{nome}</a:t></a:r></a:p></xdr:txBody></xdr:sp>')


# colunas 64 px e linhas 20 px (padrão do openpyxl)
C, L = 64, 20
anc = []
# 1. painel escuro de fundo
anc.append(ancora(1, 1, 15, 34,
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Painel de fundo"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(C, L, 14 * C, 33 * L)}<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 4000"/></a:avLst></a:prstGeom>'
    '<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="14202B"/></a:gs><a:gs pos="100000"><a:srgbClr val="23394A"/></a:gs></a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>'
    '<a:ln><a:noFill/></a:ln></xdr:spPr></xdr:sp>'))
# 2. título (caixa de texto)
anc.append(ancora(2, 2, 12, 4,
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="3" name="Título"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(2 * C, 2 * L, 10 * C, 2 * L)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></xdr:spPr>'
    + texto('PAINEL DA QUALIDADE (dados inventados)', sz=1800, b=True, algn='l', anchor='ctr') + '</xdr:sp>'))
# 3. grupo de botões de navegação
grupo = ('<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="10" name="Navegação"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>'
         f'<xdr:grpSpPr><a:xfrm><a:off x="{2 * C * EMU}" y="{5 * L * EMU}"/><a:ext cx="{9 * C * EMU}" cy="{2 * L * EMU}"/><a:chOff x="0" y="0"/><a:chExt cx="{9 * C * EMU}" cy="{2 * L * EMU}"/></a:xfrm></xdr:grpSpPr>'
         + botao(11, 'Dados', 'rId2', 0, 0, 2 * C + 40, 2 * L, sombra=True)
         + botao(12, 'Ações', 'rId3', 3 * C, 0, 2 * C + 40, 2 * L)
         + botao(13, 'Início', 'rId4', 6 * C, 0, 2 * C + 40, 2 * L) + '</xdr:grpSp>')
anc.append(ancora(2, 5, 11, 7, grupo))
# 4. seta de bloco e linha com ponta
anc.append(ancora(2, 9, 5, 11,
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="20" name="Seta Ações"><a:hlinkClick xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId3"/></xdr:cNvPr><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(2 * C, 9 * L, 3 * C, 2 * L)}<a:prstGeom prst="rightArrow"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="E0A526"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
    + texto('Ações', sz=1000, b=True, cor='1F2D36') + '</xdr:sp>'))
anc.append(ancora(5, 10, 8, 10,
    '<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="21" name="Conector"/><xdr:cNvCxnSpPr/></xdr:nvCxnSpPr>'
    f'<xdr:spPr>{xfrm(5 * C, 10 * L, 3 * C, 0)}<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="25400"><a:solidFill><a:srgbClr val="E0A526"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></xdr:spPr></xdr:cxnSp>'))
# 5. documento (fluxograma) e desenho livre
anc.append(ancora(8, 9, 11, 12,
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="22" name="Documento"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(8 * C, 9 * L, 3 * C, 3 * L)}<a:prstGeom prst="flowChartDocument"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="accent6"><a:lumMod val="50000"/></a:schemeClr></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
    + ESTILO_BOTAO + texto('POP da coleta', sz=1000, b=True) + '</xdr:sp>'))
anc.append(ancora(12, 9, 14, 12,
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="23" name="Desenho livre"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(12 * C, 9 * L, 2 * C, 3 * L)}<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst><a:path w="100" h="100">'
    '<a:moveTo><a:pt x="0" y="100"/></a:moveTo><a:lnTo><a:pt x="50" y="0"/></a:lnTo><a:cubicBezTo><a:pt x="70" y="30"/><a:pt x="90" y="60"/><a:pt x="100" y="100"/></a:cubicBezTo><a:close/></a:path></a:pathLst></a:custGeom>'
    '<a:solidFill><a:srgbClr val="3DB54A"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln></xdr:spPr></xdr:sp>'))
# 6. caixa de texto ligada à célula C40 (total)
anc.append(ancora(2, 12, 7, 14,
    '<xdr:sp macro="" textlink="$C$40"><xdr:nvSpPr><xdr:cNvPr id="24" name="Total ligado"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm(2 * C, 12 * L, 5 * C, 2 * L)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFFFFF"><a:alpha val="15000"/></a:srgbClr></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
    '<xdr:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:fld id="{1A2B3C4D-0000-0000-0000-000000000001}" type="TxLink"><a:rPr lang="pt-BR" sz="1600" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>34</a:t></a:fld></a:p></xdr:txBody></xdr:sp>'))
# 7. grupo com o gráfico dentro e uma etiqueta
gg = ('<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="30" name="Grupo do gráfico"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>'
      f'<xdr:grpSpPr><a:xfrm><a:off x="{2 * C * EMU}" y="{15 * L * EMU}"/><a:ext cx="{9 * C * EMU}" cy="{14 * L * EMU}"/><a:chOff x="0" y="0"/><a:chExt cx="{9 * C * EMU}" cy="{14 * L * EMU}"/></a:xfrm></xdr:grpSpPr>'
      '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="31" name="Etiqueta do gráfico"/><xdr:cNvSpPr/></xdr:nvSpPr>'
      f'<xdr:spPr>{xfrm(0, 0, 9 * C, 2 * L)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="2F6F8F"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
      + texto('Não conformidades (arraste o gráfico para ver)', sz=1000, b=True) + '</xdr:sp>'
      '<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="32" name="Gráfico 1"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
      f'<xdr:xfrm><a:off x="0" y="{2 * L * EMU}"/><a:ext cx="{9 * C * EMU}" cy="{12 * L * EMU}"/></xdr:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
      '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/></a:graphicData></a:graphic></xdr:graphicFrame></xdr:grpSp>')
anc.append(ancora(2, 15, 11, 29, gg))
# 8. grupo com figura (logotipo inventado) e texto
gf = ('<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="40" name="Logo"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>'
      f'<xdr:grpSpPr><a:xfrm><a:off x="{12 * C * EMU}" y="{2 * L * EMU}"/><a:ext cx="{2 * C * EMU}" cy="{3 * L * EMU}"/><a:chOff x="0" y="0"/><a:chExt cx="{2 * C * EMU}" cy="{3 * L * EMU}"/></a:xfrm></xdr:grpSpPr>'
      '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="41" name="Figura"/><xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr><xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId5"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
      f'<xdr:spPr>{xfrm(C // 2, 0, C, 2 * L)}<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic>'
      '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="42" name="Nome"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
      f'<xdr:spPr>{xfrm(0, 2 * L, 2 * C, L)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></xdr:spPr>'
      + texto('Fazenda Exemplo', sz=900, b=True, extra_bp=' lIns="0" rIns="0" tIns="0" bIns="0"') + '</xdr:sp></xdr:grpSp>')
anc.append(ancora(12, 2, 14, 5, gf))

drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + ''.join(anc) + '</xdr:wsDr>')
rels_d = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
          '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="/xl/charts/chart1.xml"/>'
          '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="#Dados!A1"/>'
          "<Relationship Id=\"rId3\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink\" Target=\"#'Ações Corretivas'!A1\"/>"
          '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="#Painel!A1"/>'
          '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>')

zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
with zipfile.ZipFile(SAIDA, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        d = zin.read(item.filename)
        if item.filename == 'xl/drawings/drawing1.xml':
            d = drawing.encode('utf-8')
        elif item.filename == 'xl/drawings/_rels/drawing1.xml.rels':
            d = rels_d.encode('utf-8')
        elif item.filename == '[Content_Types].xml':
            d = d.replace(b'<Default Extension="xml"', b'<Default Extension="png" ContentType="image/png"/><Default Extension="xml"')
        zout.writestr(item, d)
    zout.writestr('xl/media/image1.png', png(24, 24, (46, 158, 63)))
print('ok', SAIDA)
