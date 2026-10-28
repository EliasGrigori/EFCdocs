"""EFCdocs — gera testes/amostras/amostra-congelada.xlsx (DADOS INVENTADOS).
Imita o cabeçalho de um relatório: título mesclado, faixas cinzas mescladas que atravessam a linha do congelamento
(linhas 1 a 7 congeladas) e dados embaixo. Serve para o defeito da v0.6.0: ao rolar, o cinza pintava por cima do título."""
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
wb = Workbook(); ws = wb.active; ws.title = 'Relatório'
cinza = PatternFill('solid', fgColor='D9D9D9')
ws.merge_cells('B1:I4'); ws['B1'] = 'RELATÓRIO DAS ANÁLISES (EXEMPLO)'; ws['B1'].font = Font(size=18, bold=True); ws['B1'].alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('B5:F9'); ws['B5'].fill = cinza
ws.merge_cells('H5:J9'); ws['H5'].fill = cinza
ws['A7'] = 'JANEIRO 2026'; ws['A7'].font = Font(bold=True); ws['A7'].fill = cinza
ws['G7'] = 'TODOS'; ws['G7'].font = Font(bold=True); ws['G7'].fill = cinza
for r in range(10, 80):
    ws.cell(r, 1, 'Ponto %d' % (r - 9)); ws.cell(r, 2, (r * 7) % 13 / 10)
ws.freeze_panes = 'A8'
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'amostra-congelada.xlsx')); print('ok')
