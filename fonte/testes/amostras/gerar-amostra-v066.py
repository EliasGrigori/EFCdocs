"""EFCdocs — gera testes/amostras/amostra-v066.xlsx (DADOS INVENTADOS).
Imita o acabamento de um painel do Excel:
- aba Painel com largura padrão das colunas de 12,71 caracteres (89 px), um velocímetro de rosca (furo 58%) com
  título ligado à célula B2, rótulos com o "valor das células" (E2:E4), legenda e área do desenho no lugar escolhido
  à mão e uma caixa de texto por cima do gráfico ligada à célula C3; um gráfico de colunas com legenda à direita;
  linha diagonal, botões de canto recortado (snip2DiagRect e snip1Rect);
- aba Dados com uma tabela do Excel (TabColetas) e quebras de página (linha 10 e coluna C).
Uso: python3 testes/amostras/gerar-amostra-v066.py"""
import io, os, zipfile
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.pagebreak import Break

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'amostra-v066.xlsx')
EMU = 9525

wb = Workbook()
ws = wb.active
ws.title = 'Painel'
ws.sheet_format.defaultColWidth = 12.7109375
ws['B2'] = 'Cloro médio da semana'
ws['B3'] = 'Valor'; ws['C3'] = 0.62
ws['D2'] = 'Faixa'; ws['E2'] = 'Baixo'; ws['E3'] = 'Bom'; ws['E4'] = 'Alto'
ws['F2'] = 0.3; ws['F3'] = 0.7; ws['F4'] = 1.0
dados = wb.create_sheet('Dados')
dados.append(['Semana', 'Ponto', 'Cloro'])
for i, (s, p, v) in enumerate([('S1', 'Poço 1', 0.55), ('S1', 'Poço 2', 0.61), ('S2', 'Poço 1', 0.48), ('S2', 'Poço 2', 0.66),
                                ('S3', 'Poço 1', 0.71), ('S3', 'Poço 2', 0.59), ('S4', 'Poço 1', 0.63)], start=2):
    dados.cell(i, 1, s); dados.cell(i, 2, p); dados.cell(i, 3, v)
for i in range(9, 31):
    dados.cell(i, 1, f'S{i}'); dados.cell(i, 2, 'Poço 3'); dados.cell(i, 3, round(0.5 + (i % 5) * 0.05, 2))
tab = Table(displayName='TabColetas', ref='A1:C8')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9', showRowStripes=True, showColumnStripes=False)
dados.add_table(tab)
dados.row_breaks.append(Break(id=10))
dados.col_breaks.append(Break(id=3))
# dois gráficos (o 1º vira o velocímetro de rosca escrito à mão abaixo; o 2º fica de colunas)
ch1 = BarChart(); ch1.add_data(Reference(ws, min_col=6, min_row=2, max_row=4)); ws.add_chart(ch1, 'B6')
ch2 = BarChart(); ch2.title = 'Cloro por semana'
ch2.add_data(Reference(dados, min_col=3, min_row=1, max_row=8), titles_from_data=True)
ch2.set_categories(Reference(dados, min_col=1, min_row=2, max_row=8))
ch2.legend.position = 'r'
ws.add_chart(ch2, 'J6')
buf = io.BytesIO()
wb.save(buf)

NS = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
# velocímetro de rosca: título ligado a B2, rótulos com o valor das células E2:E4, legenda e área à mão, caixa por cima
pts = ''.join(f'<c:dPt><c:idx val="{i}"/><c:bubble3D val="0"/><c:spPr><a:solidFill><a:srgbClr val="{cor}"/></a:solidFill></c:spPr></c:dPt>'
              for i, cor in enumerate(['C0392B', '2E8B57', 'E0A526']))
chart1 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
          f'<c:chartSpace {NS}><c:roundedCorners val="0"/><c:chart>'
          '<c:title><c:tx><c:strRef><c:f>Painel!$B$2</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Cloro médio da semana</c:v></c:pt></c:strCache></c:strRef></c:tx>'
          '<c:layout><c:manualLayout><c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0.05"/><c:y val="0.03"/></c:manualLayout></c:layout><c:overlay val="0"/>'
          '<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="1200" b="1"/></a:pPr><a:endParaRPr lang="pt-BR"/></a:p></c:txPr></c:title>'
          '<c:autoTitleDeleted val="0"/><c:plotArea>'
          '<c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0.1"/><c:y val="0.2"/><c:w val="0.55"/><c:h val="0.75"/></c:manualLayout></c:layout>'
          '<c:doughnutChart><c:varyColors val="1"/><c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:v>Faixas</c:v></c:tx>' + pts +
          '<c:dLbls><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>'
          '<c:extLst><c:ext uri="{CE6537A1-D6FC-4f65-9D91-7224C49458BB}" xmlns:c15="http://schemas.microsoft.com/office/drawing/2012/chart"><c15:showDataLabelsRange val="1"/></c:ext></c:extLst></c:dLbls>'
          '<c:cat><c:strRef><c:f>Painel!$E$2:$E$4</c:f><c:strCache><c:ptCount val="3"/><c:pt idx="0"><c:v>Baixo</c:v></c:pt><c:pt idx="1"><c:v>Bom</c:v></c:pt><c:pt idx="2"><c:v>Alto</c:v></c:pt></c:strCache></c:strRef></c:cat>'
          '<c:val><c:numRef><c:f>Painel!$F$2:$F$4</c:f><c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="3"/><c:pt idx="0"><c:v>0.3</c:v></c:pt><c:pt idx="1"><c:v>0.7</c:v></c:pt><c:pt idx="2"><c:v>1</c:v></c:pt></c:numCache></c:numRef></c:val>'
          '<c:extLst><c:ext uri="{02D57815-91ED-43cb-92C2-25804820EDAC}" xmlns:c15="http://schemas.microsoft.com/office/drawing/2012/chart"><c15:datalabelsRange><c15:f>Painel!$E$2:$E$4</c15:f>'
          '<c15:dlblRangeCache><c:ptCount val="3"/><c:pt idx="0"><c:v>Baixo</c:v></c:pt><c:pt idx="1"><c:v>Bom</c:v></c:pt><c:pt idx="2"><c:v>Alto</c:v></c:pt></c15:dlblRangeCache></c15:datalabelsRange></c:ext></c:extLst>'
          '</c:ser><c:firstSliceAng val="270"/><c:holeSize val="58"/></c:doughnutChart></c:plotArea>'
          '<c:legend><c:legendPos val="r"/><c:layout><c:manualLayout><c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0.72"/><c:y val="0.3"/><c:w val="0.25"/><c:h val="0.4"/></c:manualLayout></c:layout><c:overlay val="0"/></c:legend>'
          '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
          '<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill></a:ln></c:spPr>'
          '<c:userShapes r:id="rId1"/></c:chartSpace>')
user_shapes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<c:userShapes xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
               '<cdr:relSizeAnchor xmlns:cdr="http://schemas.openxmlformats.org/drawingml/2006/chartDrawing">'
               '<cdr:from><cdr:x>0.25</cdr:x><cdr:y>0.5</cdr:y></cdr:from><cdr:to><cdr:x>0.5</cdr:x><cdr:y>0.68</cdr:y></cdr:to>'
               '<cdr:sp macro="" textlink="Painel!$C$3"><cdr:nvSpPr><cdr:cNvPr id="2" name="Valor do cloro"/><cdr:cNvSpPr txBox="1"/></cdr:nvSpPr>'
               '<cdr:spPr><a:xfrm xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>'
               '<a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" prst="rect"><a:avLst/></a:prstGeom><a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/></cdr:spPr>'
               '<cdr:txBody><a:bodyPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" anchor="ctr"/><a:lstStyle xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
               '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:pPr algn="ctr"/><a:fld id="{1A2B3C4D-0000-0000-0000-000000000009}" type="TxLink"><a:rPr lang="pt-BR" sz="1600" b="1"><a:solidFill><a:srgbClr val="1F2D36"/></a:solidFill></a:rPr><a:t>0,62</a:t></a:fld></a:p>'
               '</cdr:txBody></cdr:sp></cdr:relSizeAnchor></c:userShapes>')
rels_c1 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
           '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartUserShapes" Target="../drawings/drawing2.xml"/></Relationships>')

zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
nomes = zin.namelist()
# qual chart é o da aba Painel que vira velocímetro: o 1º adicionado (chart1)
desenho = zin.read('xl/drawings/drawing1.xml').decode('utf-8')
C, L = 89, 20
extra = (
    # linha diagonal (sobe da esquerda para a direita: flipV)
    '<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>24</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
    '<xdr:to><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>27</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
    '<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="50" name="Linha diagonal"/><xdr:cNvCxnSpPr/></xdr:nvCxnSpPr>'
    f'<xdr:spPr><a:xfrm flipV="1"><a:off x="{C * EMU}" y="{24 * L * EMU}"/><a:ext cx="{5 * C * EMU}" cy="{3 * L * EMU}"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
    '<a:ln w="28575"><a:solidFill><a:srgbClr val="2F6F8F"/></a:solidFill></a:ln></xdr:spPr></xdr:cxnSp><xdr:clientData/></xdr:twoCellAnchor>'
    # botão de dois cantos recortados (snip2DiagRect) e de um canto (snip1Rect)
    '<xdr:twoCellAnchor><xdr:from><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>24</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
    '<xdr:to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>27</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="51" name="Recorte diagonal"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr><a:xfrm><a:off x="{7 * C * EMU}" y="{24 * L * EMU}"/><a:ext cx="{2 * C * EMU}" cy="{3 * L * EMU}"/></a:xfrm><a:prstGeom prst="snip2DiagRect"><a:avLst><a:gd name="adj1" fmla="val 0"/><a:gd name="adj2" fmla="val 30000"/></a:avLst></a:prstGeom>'
    '<a:solidFill><a:srgbClr val="F28C28"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
    '<xdr:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="pt-BR" sz="1100" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Recorte duplo</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'
    '<xdr:twoCellAnchor><xdr:from><xdr:col>10</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>24</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
    '<xdr:to><xdr:col>12</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>27</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="52" name="Recorte simples"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr><a:xfrm><a:off x="{10 * C * EMU}" y="{24 * L * EMU}"/><a:ext cx="{2 * C * EMU}" cy="{3 * L * EMU}"/></a:xfrm><a:prstGeom prst="snip1Rect"><a:avLst/></a:prstGeom>'
    '<a:solidFill><a:srgbClr val="2E8B57"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
    '<xdr:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="pt-BR" sz="1100" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>Recorte simples</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'
)
desenho = desenho.replace('<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">',
                          '<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">', 1)
assert '</wsDr>' in desenho
desenho = desenho.replace('</wsDr>', extra + '</wsDr>')
with zipfile.ZipFile(SAIDA, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        d = zin.read(item.filename)
        if item.filename == 'xl/charts/chart1.xml':
            d = chart1.encode('utf-8')
        elif item.filename == 'xl/drawings/drawing1.xml':
            d = desenho.encode('utf-8')
        elif item.filename == '[Content_Types].xml':
            d = d.replace(b'</Types>', b'<Override PartName="/xl/drawings/drawing2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chartshapes+xml"/></Types>')
        zout.writestr(item, d)
    zout.writestr('xl/charts/_rels/chart1.xml.rels', rels_c1)
    zout.writestr('xl/drawings/drawing2.xml', user_shapes)
print('ok', SAIDA)
