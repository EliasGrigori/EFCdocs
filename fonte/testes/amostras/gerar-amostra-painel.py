"""EFCdocs — gera testes/amostras/amostra-painel.xlsx (DADOS INVENTADOS).
Imita os truques de formatação de um painel feito no Excel: velocímetro (rosca com metade invisível + pizza ponteiro),
barras só com contorno e brilho com rótulos em R$, colunas em degradê + linhas curvas com marcadores no eixo da direita
(mínimo -0,2 e máximo 1), fundo escuro, pizza com muitas fatias e colunas com eixo em R$.
Uso: python3 testes/amostras/gerar-amostra-painel.py"""
import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.fill import GradientFillProperties, GradientStop, LinearShadeProperties, ColorChoice
from openpyxl.drawing.effect import EffectList, GlowEffect
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font
from openpyxl.chart.layout import Layout, ManualLayout

AQUI = os.path.dirname(os.path.abspath(__file__))
wb = Workbook()

# ---------------- aba Verificação: velocímetros ----------------
ws = wb.active
ws.title = 'Verificação'
ws['A1'] = 'Ponto'; ws['B1'] = 'Cloro (mg/L)'
pontos = [('Cozinha', 0.58), ('Galpão', 1.35), ('Escritório', 0.0)]
for i, (p, v) in enumerate(pontos, start=2):
    ws.cell(i, 1, p); ws.cell(i, 2, v)
# faixas do velocímetro: 16 fatias iguais (0 a 2 mg/L) + metade escondida
ws['D1'] = 'Faixas'
for i in range(16):
    ws.cell(2 + i, 4, 0.125)
ws.cell(18, 4, 2.0)  # metade de baixo (invisível)
CORES = ['E2701F', '3DB54A', '2E9E3F', '1B7F2D', 'A8D43A', 'F2E33A', 'F7B32B', 'F28A1E', 'EC5A1C', 'E0251B', 'C91B1B', 'A31919', '7D1414', '5A0F0F', '3C0A0A', '000000']
# ponteiro: valor, espessura, resto
for k, (p, v) in enumerate(pontos):
    col = 6 + k
    ws.cell(1, col, 'Ponteiro ' + p)
    ws.cell(2, col, v)
    ws.cell(3, col, 0.03)
    ws.cell(4, col, 4 - v - 0.03)


def velocimetro(k, ancora):
    d = DoughnutChart()
    d.firstSliceAng = 270
    d.holeSize = 55
    d.add_data(Reference(ws, min_col=4, min_row=2, max_row=18), titles_from_data=False)
    s = d.series[0]
    for i in range(17):
        pt = DataPoint(idx=i)
        if i < 16:
            pt.graphicalProperties = GraphicalProperties(solidFill=CORES[i])
            pt.graphicalProperties.line = LineProperties(solidFill='FFFFFF', w=6350)
        else:
            pt.graphicalProperties = GraphicalProperties(noFill=True)
            pt.graphicalProperties.line = LineProperties(noFill=True)
        s.dPt.append(pt)
    p = PieChart()
    p.firstSliceAng = 270
    col = 6 + k
    p.add_data(Reference(ws, min_col=col, min_row=2, max_row=4), titles_from_data=False)
    ps = p.series[0]
    for i in range(3):
        pt = DataPoint(idx=i)
        pt.graphicalProperties = GraphicalProperties(solidFill='000000') if i == 1 else GraphicalProperties(noFill=True)
        pt.graphicalProperties.line = LineProperties(solidFill='000000', w=25400) if i == 1 else LineProperties(noFill=True)
        ps.dPt.append(pt)
    d += p
    d.legend = None
    d.title = pontos[k][0]
    d.width, d.height = 6, 4.5
    ws.add_chart(d, ancora)


velocimetro(0, 'A21'); velocimetro(1, 'E21'); velocimetro(2, 'I21')

# ---------------- aba Painel ----------------
pn = wb.create_sheet('Painel')
pn.append(['Ponto', 'Gasto (R$)', 'Coletas', 'E. coli', 'Coliformes'])
dados = [('BR 01 Bebedouro', 157.5, 0, 0, 0), ('BR 02 Bebedouro', 708.75, 1, 0, 1), ('BR 03 Bebedouro', 945.0, 1, 0, 1),
         ('BR 04 Bebedouro', 472.5, 1, 0, 0), ('BR 06 Bebedouro', 866.25, 1, 0, 0)]
for r in dados:
    pn.append(list(r))
for c in pn['B'][1:]:
    c.number_format = '"R$" #,##0.00'

ESCURO = '1F2A33'


def fundo_escuro(ch):
    ch.graphical_properties = GraphicalProperties(solidFill=ESCURO)
    ch.graphical_properties.line = LineProperties(noFill=True)
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True)


def rotulos(pos=None, fmt=None, val=True, pct=False):
    d = DataLabelList(); d.showVal = val; d.showPercent = pct; d.showCatName = False; d.showSerName = False; d.showLegendKey = False; d.showLeaderLines = False
    if pos: d.position = pos
    if fmt: d.numFmt = fmt
    return d


def texto(cor='FFFFFF', tam=900, negrito=False):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=tam, b=negrito, solidFill=cor)), endParaRPr=CharacterProperties(sz=tam, b=negrito, solidFill=cor))])


# 1) barras deitadas só com contorno e brilho, rótulos em R$
b = BarChart()
b.type = 'bar'
b.add_data(Reference(pn, min_col=2, min_row=1, max_row=6), titles_from_data=True)
b.set_categories(Reference(pn, min_col=1, min_row=2, max_row=6))
s = b.series[0]
s.graphicalProperties = GraphicalProperties(noFill=True)
s.graphicalProperties.line = LineProperties(solidFill='2E9FD8', w=19050)
s.graphicalProperties.effectLst = EffectList(glow=GlowEffect(rad=63500, srgbClr='2E9FD8'))
s.dLbls = rotulos('ctr', '"R$" #,##0.00')
s.dLbls.txPr = texto('FFFFFF', 900)
b.title = 'Gasto por ponto'
b.legend = None
b.y_axis.number_format = '"R$" #,##0.00'
b.y_axis.scaling.min = 0; b.y_axis.scaling.max = 1000
b.y_axis.txPr = texto('FFFFFF', 800); b.x_axis.txPr = texto('FFFFFF', 800, True)
b.y_axis.majorGridlines = None
b.gapWidth = 60
fundo_escuro(b)
b.width, b.height = 16, 7
pn.add_chart(b, 'G2')

# 2) colunas em degradê + linhas curvas no eixo da direita (mín -0,2, máx 1)
c = BarChart()
c.add_data(Reference(pn, min_col=3, min_row=1, max_row=6), titles_from_data=True)
c.set_categories(Reference(pn, min_col=1, min_row=2, max_row=6))
gs = GradientFillProperties(gsLst=[GradientStop(pos=0, srgbClr='1F8FC9'), GradientStop(pos=100000, srgbClr='0E3D5C')], lin=LinearShadeProperties(ang=5400000, scaled=False))
c.series[0].graphicalProperties = GraphicalProperties(gradFill=gs)
c.series[0].graphicalProperties.line = LineProperties(solidFill='0E3D5C')
c.series[0].dLbls = rotulos('inEnd'); c.series[0].dLbls.txPr = texto('FFFFFF', 1000, True)
c.gapWidth = 80
l = LineChart()
l.add_data(Reference(pn, min_col=4, max_col=5, min_row=1, max_row=6), titles_from_data=True)
for i, cor in enumerate(['8B0000', 'D2691E']):
    ls = l.series[i]
    ls.smooth = True
    ls.graphicalProperties = GraphicalProperties()
    ls.graphicalProperties.line = LineProperties(solidFill=cor, w=22225, prstDash='dash' if i == 0 else 'sysDash')
    ls.marker = Marker(symbol='circle', size=8)
    ls.marker.graphicalProperties = GraphicalProperties(solidFill=cor)
    ls.marker.graphicalProperties.line = LineProperties(solidFill='FFFFFF')
l.y_axis.axId = 200
l.y_axis.crosses = 'max'
l.y_axis.scaling.min = -0.2; l.y_axis.scaling.max = 1
l.y_axis.txPr = texto('FFFFFF', 800)
c.y_axis.txPr = texto('FFFFFF', 800); c.x_axis.txPr = texto('FFFFFF', 800, True)
c += l
c.legend.position = 'b'
c.legend.txPr = texto('FFFFFF', 900, True)
c.title = 'Análises por ponto'
fundo_escuro(c)
c.width, c.height = 16, 8
pn.add_chart(c, 'G18')

# 3) colunas com eixo em R$ e contorno sem preenchimento
g = BarChart()
pn['A9'] = 'Laboratório'; pn['B9'] = 'Gasto total'
pn['A10'] = 'Lab A'; pn['B10'] = 3638.78
pn['A11'] = 'Lab B'; pn['B11'] = 1798.0
g.add_data(Reference(pn, min_col=2, min_row=9, max_row=11), titles_from_data=True)
g.set_categories(Reference(pn, min_col=1, min_row=10, max_row=11))
g.series[0].graphicalProperties = GraphicalProperties(noFill=True)
g.series[0].graphicalProperties.line = LineProperties(solidFill='3AA0D8', w=19050)
g.series[0].dLbls = rotulos('outEnd', '"R$" #,##0.00')
g.series[0].dLbls.txPr = texto('FFFFFF', 1000, True)
g.y_axis.number_format = '"R$" #,##0'
g.y_axis.txPr = texto('FFFFFF', 800); g.x_axis.txPr = texto('FFFFFF', 800)
g.legend = None
g.title = 'Gasto total com análise'
fundo_escuro(g)
g.width, g.height = 9, 7
pn.add_chart(g, 'Q2')

# 4) pizza com muitas fatias (sem cor escrita: o Excel pinta pelas cores do tema)
pz = wb.create_sheet('Resumo')
pz.append(['Item', 'Quantidade'])
itens = [('Rede', 59), ('Bebedouro', 19), ('Reservatório', 3)] + [('Item %d' % i, 1) for i in range(1, 16)]
for it in itens:
    pz.append(list(it))
p = PieChart()
p.add_data(Reference(pz, min_col=2, min_row=1, max_row=len(itens) + 1), titles_from_data=True)
p.set_categories(Reference(pz, min_col=1, min_row=2, max_row=len(itens) + 1))
p.dataLabels = rotulos(val=False, pct=True)
p.title = 'Todos'
p.width, p.height = 12, 9
pz.add_chart(p, 'D2')

saida = os.path.join(AQUI, 'amostra-painel.xlsx')
wb.save(saida)
print('ok', saida)

# a mesma amostra com o XML dos gráficos no jeito do Excel (prefixo c:), que é como o Excel grava
import re, zipfile
saida2 = os.path.join(AQUI, 'amostra-painel-excel.xlsx')
with zipfile.ZipFile(saida) as zi, zipfile.ZipFile(saida2, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        dados = zi.read(it.filename)
        if re.match(r'xl/charts/chart\d+\.xml$', it.filename):
            x = dados.decode('utf-8')
            x = x.replace('xmlns="http://schemas.openxmlformats.org/drawingml/2006/chart"', 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"')
            x = re.sub(r'<(/?)([A-Za-z]\w*)(?=[\s/>])', r'<\1c:\2', x)
            dados = x.encode('utf-8')
        zo.writestr(it, dados)
print('ok', saida2)
